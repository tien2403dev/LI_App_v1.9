# """Bảng Alarm 23 cột, chỉnh sửa, lọc cột và xuất danh sách đang hiển thị."""
# from pathlib import Path
#
# from PyQt5.QtCore import (Qt, QAbstractTableModel, QSortFilterProxyModel,
#                           QModelIndex, pyqtSignal, QPoint, QObject, QRunnable,
#                           QThreadPool)
# from PyQt5.QtGui import QColor, QFont
# from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
#                             QTableView, QAbstractItemView, QStyledItemDelegate, QComboBox,
#                             QHeaderView, QApplication, QFileDialog, QMessageBox)
# from domain.alarm import CHECK_RESULTS, COLUMNS, STATUSES, TRACKING_FIELDS
# from ui.filter_value_dialog import FilterValueDialog
#
#
# def display_value(row, field):
#     """Định dạng phần trăm; NULL hiển thị trống, không biến thành 0%."""
#     value = row.get(field)
#     if value is None:
#         return ''
#     if field in ('yield_15', 'yield_30', 'target_15', 'target_30'):
#         return f'{value:.2f}%'
#     if field == 'alarm_date':
#         return f'{value[6:8]}/{value[4:6]}/{value[:4]}'
#     return str(value)
#
#
# class AlarmTableModel(QAbstractTableModel):
#     edit_requested = pyqtSignal(object, str, str)
#
#     def __init__(self, parent=None):
#         """Khởi tạo model nhẹ, chỉ lưu các dòng đã đọc từ database."""
#         super().__init__(parent)
#         self.rows = []
#         self.editing_enabled = True
#
#     def rowCount(self, parent=QModelIndex()):
#         """Số dòng bảng phẳng."""
#         return 0 if parent.isValid() else len(self.rows)
#
#     def columnCount(self, parent=QModelIndex()):
#         """Giữ đúng thứ tự 23 cột đã thống nhất."""
#         return 0 if parent.isValid() else len(COLUMNS)
#
#     def data(self, index, role=Qt.DisplayRole):
#         """Cấp giá trị, tooltip và màu trạng thái."""
#         if not index.isValid():
#             return None
#         row = self.rows[index.row()]
#         field = COLUMNS[index.column()][1]
#         if role in (Qt.DisplayRole, Qt.ToolTipRole):
#             return index.row() + 1 if field is None else display_value(row, field)
#         if role == Qt.EditRole:
#             return row.get(field, '')
#         # if role == Qt.BackgroundRole and field == 'status':
#         #     return QColor(dict(zip(STATUSES, ('#FDE2E2', '#FFF2CC', '#D9EAD3')))[row['status']])
#         if role == Qt.BackgroundRole:
#             status = row.get('status')
#
#             # Chưa tiến hành → cả dòng màu đỏ
#             if status == 'Chưa tiến hành':
#                 return QColor('#FDE2E2')
#
#             # Đã hoàn thành → cả dòng màu xanh
#             if status == 'Đã hoàn thành':
#                 return QColor('#D9EAD3')
#
#             # Đang tiến hành → màu vàng
#             if status == 'Đang tiến hành':
#                 return QColor('#FFF2CC')
#
#             # Quick check / Cal check
#             if field in ('quick_check_result', 'cal_check_result'):
#                 return {
#                     'PASS': QColor('#D9EAD3'),
#                     'FAIL': QColor('#FDE2E2')
#                 }.get(str(row.get(field, '')).upper())
#
#             return None
#
#         if role == Qt.TextAlignmentRole:
#             return int(Qt.AlignVCenter | Qt.AlignHCenter)
#         return None
#
#     def headerData(self, section, orientation, role=Qt.DisplayRole):
#         """Tên cột; proxy thêm dấu lọc."""
#         if orientation == Qt.Horizontal and role == Qt.DisplayRole:
#             return COLUMNS[section][0]
#         return super().headerData(section, orientation, role)
#
#     def flags(self, index):
#         """Chỉ cho sửa Status và các trường kết quả xử lý."""
#         flags = super().flags(index)
#         if index.isValid() and self.editing_enabled and COLUMNS[index.column()][1] in TRACKING_FIELDS:
#             flags |= Qt.ItemIsEditable
#         return flags
#
#     def setData(self, index, value, role=Qt.EditRole):
#         """Gửi yêu cầu lưu; không đổi dữ liệu hiển thị trước commit."""
#         if not index.isValid() or role != Qt.EditRole or not self.editing_enabled:
#             return False
#         field = COLUMNS[index.column()][1]
#         if field not in TRACKING_FIELDS:
#             return False
#         row = self.rows[index.row()]
#         if str(value) != row[field]:
#             self.edit_requested.emit(dict(row), field, str(value))
#         return True
#
#     def set_rows(self, rows):
#         """Thay snapshot sau khi worker tải thành công."""
#         self.beginResetModel()
#         self.rows = rows
#         self.endResetModel()
#
#     def update_row(self, latest):
#         """Thay đúng ID, không dựa vào vị trí sau lọc/sắp xếp."""
#         for index, row in enumerate(self.rows):
#             if row['id'] == latest['id']:
#                 self.rows[index] = latest
#                 self.dataChanged.emit(self.index(index, 0), self.index(index, len(COLUMNS)-1))
#                 break
#
#
# class AlarmFilterProxy(QSortFilterProxyModel):
#     def __init__(self, parent=None):
#         """Lọc AND giữa các cột; không thay dữ liệu được dùng để tạo alarm."""
#         super().__init__(parent)
#         self.filters = {}
#
#     def filterAcceptsRow(self, source_row, source_parent):
#         """Một dòng phải đạt tất cả lựa chọn của các cột."""
#         model = self.sourceModel()
#         return all(str(model.data(model.index(source_row, col))) in selected
#                    for col, selected in self.filters.items())
#
#     def data(self, index, role=Qt.DisplayRole):
#         """Đánh số No lại theo thứ tự đang hiển thị."""
#         if index.isValid() and index.column() == 0 and role == Qt.DisplayRole:
#             return index.row() + 1
#         return super().data(index, role)
#
#     def headerData(self, section, orientation, role=Qt.DisplayRole):
#         """Đánh dấu rõ cột đang có bộ lọc."""
#         value = super().headerData(section, orientation, role)
#         if orientation == Qt.Horizontal and role == Qt.DisplayRole and section:
#             return f'{value} {"●" if section in self.filters else "▾"}'
#         return value
#
#     def set_filter(self, column, selected, values):
#         """Chọn tất cả thì bỏ điều kiện; chọn rỗng thì không có dòng nào."""
#         if selected == set(values):
#             self.filters.pop(column, None)
#         else:
#             self.filters[column] = set(selected)
#         self.invalidateFilter()
#         self.headerDataChanged.emit(Qt.Horizontal, column, column)
#
#     def clear_filters(self):
#         """Xóa điều kiện cột và cập nhật biểu tượng."""
#         self.filters.clear()
#         self.invalidateFilter()
#         self.headerDataChanged.emit(Qt.Horizontal, 0, len(COLUMNS)-1)
#
#
# class StatusDelegate(QStyledItemDelegate):
#     def createEditor(self, parent, option, index):
#         """Dropdown chỉ chứa ba trạng thái được phép."""
#         editor = QComboBox(parent)
#         editor.addItems(STATUSES)
#         editor.activated.connect(lambda _index: self._commit(editor))
#         return editor
#
#     def setEditorData(self, editor, index):
#         """Chọn trạng thái đã lưu trước khi người dùng chỉnh."""
#         editor.setCurrentText(index.data(Qt.EditRole))
#
#     def setModelData(self, editor, model, index):
#         """Chuyển giá trị sang yêu cầu lưu của model."""
#         model.setData(index, editor.currentText(), Qt.EditRole)
#
#     def _commit(self, editor):
#         """Lưu ngay khi người dùng chọn và đóng editor."""
#         self.commitData.emit(editor)
#         self.closeEditor.emit(editor)
#
#
# class CheckResultDelegate(QStyledItemDelegate):
#     """Chỉ cho chọn PASS/FAIL, không tạo ô nhập chữ tự do."""
#
#     def createEditor(self, parent, option, index):
#         """Tạo combobox không cho phép nhập nội dung ngoài danh sách."""
#         editor = QComboBox(parent)
#         editor.setEditable(False)
#         editor.addItems(CHECK_RESULTS)
#         editor.activated.connect(lambda _index: self._commit(editor))
#         return editor
#
#     def setEditorData(self, editor, index):
#         """Hiển thị giá trị đã lưu; alarm mới có thể đang để trống."""
#         current = str(index.data(Qt.EditRole) or '').upper()
#         position = editor.findText(current)
#         editor.setCurrentIndex(position if position >= 0 else -1)
#
#     def setModelData(self, editor, model, index):
#         """Chỉ gửi PASS/FAIL được chọn tới cơ chế lưu hiện tại."""
#         value = editor.currentText()
#         if value in CHECK_RESULTS:
#             model.setData(index, value, Qt.EditRole)
#
#     def _commit(self, editor):
#         """Lưu ngay khi chọn và đóng combobox."""
#         self.commitData.emit(editor)
#         self.closeEditor.emit(editor)
#
#
# class AlarmExportSignals(QObject):
#     """Trả kết quả ghi Excel từ worker về UI thread."""
#
#     finished = pyqtSignal(str, str)
#
#
# class AlarmExportWorker(QRunnable):
#     """Ghi snapshot các dòng alarm đang hiển thị mà không chặn giao diện."""
#
#     def __init__(self, output_path, headers, rows):
#         """Giữ dữ liệu thuần, không truy cập widget từ luồng nền."""
#         super().__init__()
#         self.output_path = output_path
#         self.headers = headers
#         self.rows = rows
#         self.signals = AlarmExportSignals()
#
#     def run(self):
#         """Tạo workbook theo định dạng Export Alarm của Aging."""
#         workbook = None
#         error_message = ''
#         try:
#             # Chỉ tải openpyxl khi người dùng thực sự bấm Export Excel.
#             from openpyxl import Workbook
#             from openpyxl.styles import Alignment, Font, PatternFill
#             from openpyxl.utils import get_column_letter
#
#             workbook = Workbook()
#             sheet = workbook.active
#             sheet.title = 'List of Alarm'
#             sheet.append(self.headers)
#
#             for values in self.rows:
#                 sheet.append(values)
#                 # Mã hoặc ghi chú bắt đầu bằng "=" phải được giữ như text.
#                 for cell in sheet[sheet.max_row]:
#                     if isinstance(cell.value, str):
#                         cell.data_type = 's'
#
#             header_fill = PatternFill(fill_type='solid', fgColor='E2E8F0')
#             for cell in sheet[1]:
#                 cell.font = Font(bold=True, color='1E293B')
#                 cell.fill = header_fill
#                 cell.alignment = Alignment(horizontal='center', vertical='center')
#             sheet.row_dimensions[1].height = 26
#
#             for row in sheet.iter_rows(min_row=2):
#                 for cell in row:
#                     cell.alignment = Alignment(horizontal='center', vertical='center')
#
#             last_column = get_column_letter(len(self.headers))
#             sheet.auto_filter.ref = f'A1:{last_column}{sheet.max_row}'
#             sheet.freeze_panes = 'A2'
#
#             width_by_header = {
#                 'No': 7, 'Alarm Date': 14, 'EQP': 12, 'Slot': 9,
#                 'Fail comment': 48, 'Start time': 20, 'End time': 20,
#                 'Yield 15': 12, 'Target 15': 12,
#                 'Yield 30': 12, 'Target 30': 12, 'Scrap code': 18,
#                 'LOTID': 28, 'Model': 13, 'Status': 16,
#                 'Date complete': 22, 'Quick check result': 28,
#                 'Cal check result': 28, 'Engineer action': 32,
#                 'Monitor day 1': 24, 'Monitor day 2': 24,
#                 'Monitor day 3': 24, 'Comment': 36,
#             }
#             for column, header in enumerate(self.headers, start=1):
#                 width = width_by_header.get(header, 18)
#                 sheet.column_dimensions[get_column_letter(column)].width = width
#
#             workbook.save(self.output_path)
#         except Exception as error:
#             error_message = str(error).strip() or 'Không thể xuất danh sách alarm.'
#         finally:
#             if workbook is not None:
#                 workbook.close()
#
#         self.signals.finished.emit(self.output_path, error_message)
#
#
# class AlarmPage(QWidget):
#     history_requested = pyqtSignal(object)
#     edit_requested = pyqtSignal(object, str, str)
#
#     def __init__(self, parent=None):
#         """Tạo bảng lazily, không đọc DB hoặc import thư viện nặng."""
#         super().__init__(parent)
#         layout = QVBoxLayout(self)
#         title = QLabel('List of Alarm')
#         title.setStyleSheet('font-size:18px; font-weight:600; color:#245585;')
#         layout.addWidget(title)
#         bar = QHBoxLayout()
#         self.status_label = QLabel('Chọn From / To rồi bấm Search.')
#         bar.addWidget(self.status_label, 1)
#         self.clear_button = QPushButton('Clear column filters')
#         bar.addWidget(self.clear_button)
#         self.export_button = QPushButton('Export Excel')
#         self.export_button.setMinimumWidth(130)
#         self.export_button.setMinimumHeight(30)
#         self.export_button.setStyleSheet('''
#             QPushButton {background:#1976D2; color:white; border:none;
#                          border-radius:4px; padding:6px 14px; font-weight:600;}
#             QPushButton:hover {background:#1565C0;}
#             QPushButton:pressed {background:#0D47A1;}
#             QPushButton:disabled {background:#90CAF9; color:#E3F2FD;}
#         ''')
#         self.export_button.setEnabled(False)
#         bar.addWidget(self.export_button)
#         layout.addLayout(bar)
#         self.table = QTableView()
#         self.table.setFont(QFont("Segoe UI", 10))
#         self.model = AlarmTableModel(self)
#         self.proxy = AlarmFilterProxy(self)
#         self.proxy.setSourceModel(self.model)
#         self.table.setModel(self.proxy)
#         self.table.setAlternatingRowColors(True)
#         self.table.setWordWrap(False)
#         self.table.setSelectionBehavior(QAbstractItemView.SelectItems)
#         self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
#         self.table.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed)
#         self.table.verticalHeader().hide()
#         self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
#         self.table.horizontalHeader().setSectionsClickable(True)
#         self.table.setItemDelegateForColumn(14, StatusDelegate(self.table))
#         self.table.setItemDelegateForColumn(16, CheckResultDelegate(self.table))
#         self.table.setItemDelegateForColumn(17, CheckResultDelegate(self.table))
#         widths = [48,110,100,65,350,145,145,105,82,105,82,130,190,105,115,170,
#                   200,200,230,115,115,115,250]
#         for col, width in enumerate(widths):
#             metrics = self.table.fontMetrics()
#             minimum = metrics.horizontalAdvance(COLUMNS[col][0] + ' ▾') + 24
#             if col in (5, 6, 15):
#                 minimum = max(minimum, metrics.horizontalAdvance('2026-09-15 23:59:59') + 12)
#             self.table.setColumnWidth(col, max(width, minimum))
#         # Bảng Alarm hiển thị Yield và ẩn hai giá trị Target cấu hình.
#         self.table.setColumnHidden(8, True)
#         self.table.setColumnHidden(10, True)
#         self.table.setStyleSheet('QHeaderView::section {background:#DCEAF7; padding:6px; border:1px solid #B8CDE0;}')
#         layout.addWidget(self.table, 1)
#         self.model.edit_requested.connect(self.edit_requested)
#         self.table.clicked.connect(self.open_history)
#         self.table.horizontalHeader().sectionClicked.connect(self.open_filter)
#         self.clear_button.clicked.connect(self.clear_filters)
#         self.export_button.clicked.connect(self.export_alarm_excel)
#         self.proxy.rowsInserted.connect(self.update_count)
#         self.proxy.rowsRemoved.connect(self.update_count)
#         self._busy = False
#         self._exporting = False
#         self._export_worker = None
#
#     def open_history(self, index):
#         """Chỉ nhóm cột máy (No..Model); map proxy để đúng ID sau lọc/sort."""
#         if self._busy or not index.isValid():
#             return
#         source = self.proxy.mapToSource(index)
#         # Status, completion, kiểm tra, action, monitor và comment không mở popup.
#         field = COLUMNS[source.column()][1]
#         if field in TRACKING_FIELDS or source.column() >= 14:
#             return
#         row = self.model.rows[source.row()]
#         preferred = 15 if field in ('yield_15', 'target_15') else (
#             30 if field in ('yield_30', 'target_30') else None)
#         self.history_requested.emit(dict(id=row['id'], row_version=row['row_version'],
#                                          preferred_rule=preferred))
#
#     def set_busy(self, busy):
#         """Không cho mở editor hoặc popup trong lúc đọc/lưu/import."""
#         self._busy = bool(busy)
#         self.model.editing_enabled = not busy
#         self.table.setEnabled(not busy)
#         self.clear_button.setEnabled(not busy)
#         self.export_button.setEnabled(
#             not busy and not self._exporting and self.proxy.rowCount() > 0
#         )
#
#     def load_rows(self, rows):
#         """Hiển thị snapshot mới; giữ bộ lọc cột đang chọn."""
#         self.model.set_rows(rows)
#         self.proxy.invalidateFilter()
#         self.update_count()
#
#     def update_count(self, *_args):
#         """Hiển thị số dòng sau lọc và tổng số alarm trong khoảng ngày."""
#         self.status_label.setText(f'Alarm: {self.proxy.rowCount()} / {len(self.model.rows)} | Lọc theo Alarm Date (From / To)')
#         self.export_button.setEnabled(
#             not self._busy and not self._exporting and self.proxy.rowCount() > 0
#         )
#
#     def clear_filters(self):
#         """Trả lại toàn bộ alarm đã tải trong khoảng ngày."""
#         self.proxy.clear_filters()
#         self.update_count()
#
#     def open_filter(self, column):
#         """Popup dưới tiêu đề cột, click áp dụng ngay và chỉ có Close."""
#         if column == 0 or not self.model.editing_enabled:
#             return
#         values = sorted({str(self.model.data(self.model.index(i, column)))
#                          for i in range(self.model.rowCount())})
#         selected = self.proxy.filters.get(column, set(values))
#         dialog = FilterValueDialog(COLUMNS[column][0], values, selected,
#                                    apply_immediately=True, compact=True, parent=self)
#         dialog.selection_changed.connect(lambda selected: self.apply_filter(column, selected, values))
#         header = self.table.horizontalHeader()
#         point = header.mapToGlobal(QPoint(header.sectionViewportPosition(column), header.height()))
#         screen = QApplication.screenAt(point) or QApplication.primaryScreen()
#         rect = screen.availableGeometry()
#         point.setX(max(rect.left(), min(point.x(), rect.right()-dialog.width())))
#         point.setY(max(rect.top(), min(point.y(), rect.bottom()-dialog.height())))
#         dialog.move(point)
#         dialog.exec_()
#
#     def apply_filter(self, column, selected, values):
#         """Áp dụng lựa chọn và cập nhật bộ đếm."""
#         self.proxy.set_filter(column, selected, values)
#         self.update_count()
#
#     def export_alarm_excel(self):
#         """Chụp và xuất đúng các dòng hiện thấy sau tất cả bộ lọc cột."""
#         if self._busy or self._exporting:
#             return
#         if self.proxy.rowCount() == 0:
#             QMessageBox.information(self, 'Export Excel', 'Không có alarm để xuất.')
#             return
#
#         visible_columns = [column for column in range(self.proxy.columnCount())
#                            if not self.table.isColumnHidden(column)]
#         rows = [[self.proxy.data(self.proxy.index(row, column), Qt.DisplayRole)
#                  for column in visible_columns]
#                 for row in range(self.proxy.rowCount())]
#         headers = [COLUMNS[column][0] for column in visible_columns]
#
#         alarm_dates = sorted({self.proxy.data(self.proxy.index(row, 1), Qt.EditRole)
#                               for row in range(self.proxy.rowCount())})
#         date_from = str(alarm_dates[0]).replace('/', '-') if alarm_dates else ''
#         date_to = str(alarm_dates[-1]).replace('/', '-') if alarm_dates else ''
#         default_name = f'List of Alarm from {date_from} - {date_to}.xlsx'
#
#         output_path, _ = QFileDialog.getSaveFileName(
#             self, 'Export danh sách Alarm', default_name, 'Excel Workbook (*.xlsx)')
#         if not output_path:
#             return
#         if not output_path.lower().endswith('.xlsx'):
#             output_path += '.xlsx'
#             if Path(output_path).exists():
#                 answer = QMessageBox.question(
#                     self, 'Export Excel',
#                     'File đã tồn tại. Bạn có muốn ghi đè không?',
#                     QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
#                 if answer != QMessageBox.Yes:
#                     return
#
#         self._exporting = True
#         self.export_button.setEnabled(False)
#         self.export_button.setText('Exporting...')
#         self._export_worker = AlarmExportWorker(output_path, headers, rows)
#         self._export_worker.signals.finished.connect(self._on_alarm_export_finished)
#         QThreadPool.globalInstance().start(self._export_worker)
#
#     def _on_alarm_export_finished(self, output_path, error_message):
#         """Khôi phục nút và thông báo kết quả sau khi worker hoàn thành."""
#         self._exporting = False
#         self._export_worker = None
#         self.export_button.setText('Export Excel')
#         self.export_button.setEnabled(not self._busy and self.proxy.rowCount() > 0)
#         if error_message:
#             QMessageBox.warning(
#                 self, 'Export Excel', f'Xuất Excel thất bại:\n{error_message}')
#             return
#         QMessageBox.information(
#             self, 'Export Excel', f'Đã xuất danh sách alarm:\n{output_path}')
#


"""Bảng Alarm 23 cột, chỉnh sửa, lọc cột và xuất danh sách đang hiển thị."""
from pathlib import Path

from PyQt5.QtCore import (Qt, QAbstractTableModel, QSortFilterProxyModel,
                          QModelIndex, pyqtSignal, QPoint, QObject, QRunnable,
                          QThreadPool)
from PyQt5.QtGui import QColor, QFont, QPen, QPolygon
from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                             QTableView, QAbstractItemView, QStyledItemDelegate, QComboBox,
                             QHeaderView, QApplication, QFileDialog, QMessageBox)
from domain.alarm import CHECK_RESULTS, COLUMNS, STATUSES, TRACKING_FIELDS
from ui.filter_value_dialog import FilterValueDialog


def display_value(row, field):
    """Định dạng phần trăm; NULL hiển thị trống, không biến thành 0%."""
    value = row.get(field)
    if value is None:
        return ''
    if field in ('yield_15', 'yield_30', 'target_15', 'target_30'):
        return f'{value:.2f}%'
    if field == 'alarm_date':
        return f'{value[6:8]}/{value[4:6]}/{value[:4]}'
    return str(value)


class AlarmTableModel(QAbstractTableModel):
    edit_requested = pyqtSignal(object, str, str)

    def __init__(self, parent=None):
        """Khởi tạo model nhẹ, chỉ lưu các dòng đã đọc từ database."""
        super().__init__(parent)
        self.rows = []
        self.editing_enabled = True

    def rowCount(self, parent=QModelIndex()):
        """Số dòng bảng phẳng."""
        return 0 if parent.isValid() else len(self.rows)

    def columnCount(self, parent=QModelIndex()):
        """Giữ đúng thứ tự 23 cột đã thống nhất."""
        return 0 if parent.isValid() else len(COLUMNS)

    def data(self, index, role=Qt.DisplayRole):
        """Cấp giá trị, tooltip và màu trạng thái."""
        if not index.isValid():
            return None
        row = self.rows[index.row()]
        field = COLUMNS[index.column()][1]
        if role in (Qt.DisplayRole, Qt.ToolTipRole):
            return index.row() + 1 if field is None else display_value(row, field)
        if role == Qt.EditRole:
            return row.get(field, '')
        # if role == Qt.BackgroundRole and field == 'status':
        #     return QColor(dict(zip(STATUSES, ('#FDE2E2', '#FFF2CC', '#D9EAD3')))[row['status']])
        if role == Qt.BackgroundRole:
            status = row.get('status')

            # Chưa tiến hành → cả dòng màu đỏ
            if status == 'Chưa tiến hành':
                return QColor('#FDE2E2')

            # Đã hoàn thành → cả dòng màu xanh
            if status == 'Đã hoàn thành':
                return QColor('#D9EAD3')

            # Đang tiến hành → màu vàng
            if status == 'Đang tiến hành':
                return QColor('#FFF2CC')

            # Quick check / Cal check
            if field in ('quick_check_result', 'cal_check_result'):
                return {
                    'PASS': QColor('#D9EAD3'),
                    'FAIL': QColor('#FDE2E2')
                }.get(str(row.get(field, '')).upper())

            return None

        if role == Qt.TextAlignmentRole:
            return int(Qt.AlignVCenter | Qt.AlignHCenter)
        return None

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        """Tên cột; proxy thêm dấu lọc."""
        if orientation == Qt.Horizontal and role == Qt.DisplayRole:
            return COLUMNS[section][0]
        return super().headerData(section, orientation, role)

    def flags(self, index):
        """Chỉ cho sửa Status và các trường kết quả xử lý."""
        flags = super().flags(index)
        if index.isValid() and self.editing_enabled and COLUMNS[index.column()][1] in TRACKING_FIELDS:
            flags |= Qt.ItemIsEditable
        return flags

    def setData(self, index, value, role=Qt.EditRole):
        """Gửi yêu cầu lưu; không đổi dữ liệu hiển thị trước commit."""
        if not index.isValid() or role != Qt.EditRole or not self.editing_enabled:
            return False
        field = COLUMNS[index.column()][1]
        if field not in TRACKING_FIELDS:
            return False
        row = self.rows[index.row()]
        if str(value) != row[field]:
            self.edit_requested.emit(dict(row), field, str(value))
        return True

    def set_rows(self, rows):
        """Thay snapshot sau khi worker tải thành công."""
        self.beginResetModel()
        self.rows = rows
        self.endResetModel()

    def update_row(self, latest):
        """Thay đúng ID, không dựa vào vị trí sau lọc/sắp xếp."""
        for index, row in enumerate(self.rows):
            if row['id'] == latest['id']:
                self.rows[index] = latest
                self.dataChanged.emit(self.index(index, 0), self.index(index, len(COLUMNS) - 1))
                break


class AlarmFilterProxy(QSortFilterProxyModel):
    def __init__(self, parent=None):
        """Lọc AND giữa các cột; không thay dữ liệu được dùng để tạo alarm."""
        super().__init__(parent)
        self.filters = {}

    def filterAcceptsRow(self, source_row, source_parent):
        """Một dòng phải đạt tất cả lựa chọn của các cột."""
        model = self.sourceModel()
        return all(str(model.data(model.index(source_row, col))) in selected
                   for col, selected in self.filters.items())

    def data(self, index, role=Qt.DisplayRole):
        """Đánh số No lại theo thứ tự đang hiển thị."""
        if index.isValid() and index.column() == 0 and role == Qt.DisplayRole:
            return index.row() + 1
        return super().data(index, role)

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        """Tên cột giữ nguyên; HeaderView tự vẽ trạng thái bộ lọc."""
        value = super().headerData(section, orientation, role)
        if orientation == Qt.Horizontal and role == Qt.ToolTipRole:
            if section in self.filters:
                return 'Đang lọc cột này. Bấm để thay đổi bộ lọc.'
            return 'Bấm để lọc cột này.'
        return value

    def set_filter(self, column, selected, values):
        """Chọn tất cả thì bỏ điều kiện; chọn rỗng thì không có dòng nào."""
        if selected == set(values):
            self.filters.pop(column, None)
        else:
            self.filters[column] = set(selected)
        self.invalidateFilter()
        self.headerDataChanged.emit(Qt.Horizontal, column, column)

    def clear_filters(self):
        """Xóa điều kiện cột và cập nhật biểu tượng."""
        self.filters.clear()
        self.invalidateFilter()
        self.headerDataChanged.emit(Qt.Horizontal, 0, len(COLUMNS) - 1)

    def is_column_filtered(self, column):
        """Cho HeaderView biết cột nào đang có điều kiện lọc."""
        return column in self.filters


class AlarmFilterHeader(QHeaderView):
    """Vẽ mũi tên và trạng thái lọc giống bảng Alarm của Aging."""

    def paintSection(self, painter, rect, logical_index):
        if not rect.isValid():
            return

        model = self.model()
        is_filtered = (
                isinstance(model, AlarmFilterProxy)
                and model.is_column_filtered(logical_index)
        )

        painter.save()
        background = QColor('#DBEAFE' if is_filtered else '#E2E8F0')
        foreground = QColor('#0D47A1' if is_filtered else '#1E293B')
        painter.fillRect(rect, background)
        painter.setPen(QPen(QColor('#CBD5E1'), 1))
        painter.drawLine(rect.topRight(), rect.bottomRight())
        painter.drawLine(rect.bottomLeft(), rect.bottomRight())

        font = painter.font()
        font.setBold(True)
        painter.setFont(font)
        painter.setPen(foreground)
        header_text = model.headerData(
            logical_index, Qt.Horizontal, Qt.DisplayRole
        )
        text_rect = rect.adjusted(4, 0, -20, 0)
        painter.drawText(
            text_rect,
            int(Qt.AlignCenter | Qt.AlignVCenter | Qt.TextSingleLine),
            str(header_text or ''),
        )

        icon_x = rect.right() - 10
        icon_y = rect.center().y()
        painter.setPen(Qt.NoPen)
        if is_filtered:
            painter.setBrush(QColor('#1565C0'))
            funnel = QPolygon([
                QPoint(icon_x - 6, icon_y - 5),
                QPoint(icon_x + 6, icon_y - 5),
                QPoint(icon_x + 2, icon_y),
                QPoint(icon_x + 2, icon_y + 5),
                QPoint(icon_x - 2, icon_y + 3),
                QPoint(icon_x - 2, icon_y),
            ])
            painter.drawPolygon(funnel)
            painter.fillRect(
                rect.left(), rect.bottom() - 2, rect.width(), 3,
                QColor('#1976D2'),
            )
        else:
            painter.setBrush(QColor('#334155'))
            arrow = QPolygon([
                QPoint(icon_x - 5, icon_y - 3),
                QPoint(icon_x + 5, icon_y - 3),
                QPoint(icon_x, icon_y + 3),
            ])
            painter.drawPolygon(arrow)
        painter.restore()


class StatusDelegate(QStyledItemDelegate):
    def createEditor(self, parent, option, index):
        """Dropdown chỉ chứa ba trạng thái được phép."""
        editor = QComboBox(parent)
        editor.addItems(STATUSES)
        editor.activated.connect(lambda _index: self._commit(editor))
        return editor

    def setEditorData(self, editor, index):
        """Chọn trạng thái đã lưu trước khi người dùng chỉnh."""
        editor.setCurrentText(index.data(Qt.EditRole))

    def setModelData(self, editor, model, index):
        """Chuyển giá trị sang yêu cầu lưu của model."""
        model.setData(index, editor.currentText(), Qt.EditRole)

    def _commit(self, editor):
        """Lưu ngay khi người dùng chọn và đóng editor."""
        self.commitData.emit(editor)
        self.closeEditor.emit(editor)


class CheckResultDelegate(QStyledItemDelegate):
    """Chỉ cho chọn PASS/FAIL, không tạo ô nhập chữ tự do."""

    def createEditor(self, parent, option, index):
        """Tạo combobox không cho phép nhập nội dung ngoài danh sách."""
        editor = QComboBox(parent)
        editor.setEditable(False)
        editor.addItems(CHECK_RESULTS)
        editor.activated.connect(lambda _index: self._commit(editor))
        return editor

    def setEditorData(self, editor, index):
        """Hiển thị giá trị đã lưu; alarm mới có thể đang để trống."""
        current = str(index.data(Qt.EditRole) or '').upper()
        position = editor.findText(current)
        editor.setCurrentIndex(position if position >= 0 else -1)

    def setModelData(self, editor, model, index):
        """Chỉ gửi PASS/FAIL được chọn tới cơ chế lưu hiện tại."""
        value = editor.currentText()
        if value in CHECK_RESULTS:
            model.setData(index, value, Qt.EditRole)

    def _commit(self, editor):
        """Lưu ngay khi chọn và đóng combobox."""
        self.commitData.emit(editor)
        self.closeEditor.emit(editor)


class AlarmExportSignals(QObject):
    """Trả kết quả ghi Excel từ worker về UI thread."""

    finished = pyqtSignal(str, str)


class AlarmExportWorker(QRunnable):
    """Ghi snapshot các dòng alarm đang hiển thị mà không chặn giao diện."""

    def __init__(self, output_path, headers, rows):
        """Giữ dữ liệu thuần, không truy cập widget từ luồng nền."""
        super().__init__()
        self.output_path = output_path
        self.headers = headers
        self.rows = rows
        self.signals = AlarmExportSignals()

    def run(self):
        """Tạo workbook theo định dạng Export Alarm của Aging."""
        workbook = None
        error_message = ''
        try:
            # Chỉ tải openpyxl khi người dùng thực sự bấm Export Excel.
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'List of Alarm'
            sheet.append(self.headers)

            for values in self.rows:
                sheet.append(values)
                # Mã hoặc ghi chú bắt đầu bằng "=" phải được giữ như text.
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'

            header_fill = PatternFill(fill_type='solid', fgColor='E2E8F0')
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='1E293B')
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
            sheet.row_dimensions[1].height = 26

            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(horizontal='center', vertical='center')

            last_column = get_column_letter(len(self.headers))
            sheet.auto_filter.ref = f'A1:{last_column}{sheet.max_row}'
            sheet.freeze_panes = 'A2'

            width_by_header = {
                'No': 7, 'Alarm Date': 14, 'EQP': 12, 'Slot': 9,
                'Fail comment': 48, 'Start time': 20, 'End time': 20,
                'Yield 15': 12, 'Target 15': 12,
                'Yield 30': 12, 'Target 30': 12, 'Scrap code': 18,
                'LOTID': 28, 'Model': 13, 'Status': 16,
                'Date complete': 22, 'Quick check result': 28,
                'Cal check result': 28, 'Engineer action': 32,
                'Monitor day 1': 24, 'Monitor day 2': 24,
                'Monitor day 3': 24, 'Comment': 36,
            }
            for column, header in enumerate(self.headers, start=1):
                width = width_by_header.get(header, 18)
                sheet.column_dimensions[get_column_letter(column)].width = width

            workbook.save(self.output_path)
        except Exception as error:
            error_message = str(error).strip() or 'Không thể xuất danh sách alarm.'
        finally:
            if workbook is not None:
                workbook.close()

        self.signals.finished.emit(self.output_path, error_message)


class AlarmPage(QWidget):
    history_requested = pyqtSignal(object)
    edit_requested = pyqtSignal(object, str, str)
    dashboard_rows_changed = pyqtSignal(object)

    def __init__(self, parent=None):
        """Tạo bảng lazily, không đọc DB hoặc import thư viện nặng."""
        super().__init__(parent)
        layout = QVBoxLayout(self)
        self.title_label = QLabel('List of Alarm')
        self.title_label.setStyleSheet('font-size:18px; font-weight:600; color:#245585;')
        layout.addWidget(self.title_label)
        bar = QHBoxLayout()
        self.status_label = QLabel('Chọn From / To rồi bấm Search.')
        bar.addWidget(self.status_label, 1)
        self.export_button = QPushButton('Export Excel')
        self.export_button.setMinimumWidth(130)
        self.export_button.setMinimumHeight(30)
        self.export_button.setStyleSheet('''
            QPushButton {background:#1976D2; color:white; border:none;
                         border-radius:4px; padding:6px 14px; font-weight:600;}
            QPushButton:hover {background:#1565C0;}
            QPushButton:pressed {background:#0D47A1;}
            QPushButton:disabled {background:#90CAF9; color:#E3F2FD;}
        ''')
        self.export_button.setEnabled(False)
        bar.addWidget(self.export_button)
        layout.addLayout(bar)
        self.table = QTableView()
        self.table.setFont(QFont("Segoe UI", 10))
        self.model = AlarmTableModel(self)
        self.proxy = AlarmFilterProxy(self)
        self.proxy.setSourceModel(self.model)
        self.filter_header = AlarmFilterHeader(Qt.Horizontal, self.table)
        self.table.setHorizontalHeader(self.filter_header)
        self.table.setModel(self.proxy)
        self.table.setAlternatingRowColors(True)
        self.table.setWordWrap(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectItems)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed)
        self.table.verticalHeader().hide()
        self.filter_header.setSectionResizeMode(QHeaderView.Interactive)
        self.filter_header.setSectionsClickable(True)
        self.filter_header.setHighlightSections(False)
        self.filter_header.setFixedHeight(32)
        self.table.setItemDelegateForColumn(14, StatusDelegate(self.table))
        self.table.setItemDelegateForColumn(16, CheckResultDelegate(self.table))
        self.table.setItemDelegateForColumn(17, CheckResultDelegate(self.table))
        widths = [48, 110, 100, 65, 350, 145, 145, 105, 82, 105, 82, 130, 190, 105, 115, 170,
                  200, 200, 230, 115, 115, 115, 250]
        for col, width in enumerate(widths):
            metrics = self.table.fontMetrics()
            minimum = metrics.horizontalAdvance(COLUMNS[col][0] + ' ▾') + 24
            if col in (5, 6, 15):
                minimum = max(minimum, metrics.horizontalAdvance('2026-09-15 23:59:59') + 12)
            self.table.setColumnWidth(col, max(width, minimum))
        # Bảng Alarm hiển thị Yield và ẩn hai giá trị Target cấu hình.
        self.table.setColumnHidden(8, True)
        self.table.setColumnHidden(10, True)
        self.table.setStyleSheet('QHeaderView::section {background:#DCEAF7; padding:6px; border:1px solid #B8CDE0;}')
        layout.addWidget(self.table, 1)
        self.model.edit_requested.connect(self.edit_requested)
        self.table.clicked.connect(self.open_history)
        self.table.horizontalHeader().sectionClicked.connect(self.open_filter)
        self.export_button.clicked.connect(self.export_alarm_excel)
        self.proxy.rowsInserted.connect(self.update_count)
        self.proxy.rowsRemoved.connect(self.update_count)
        self._busy = False
        self._exporting = False
        self._export_worker = None

    def open_history(self, index):
        """Chỉ nhóm cột máy (No..Model); map proxy để đúng ID sau lọc/sort."""
        if self._busy or not index.isValid():
            return
        source = self.proxy.mapToSource(index)
        # Status, completion, kiểm tra, action, monitor và comment không mở popup.
        field = COLUMNS[source.column()][1]
        if field in TRACKING_FIELDS or source.column() >= 14:
            return
        row = self.model.rows[source.row()]
        preferred = 15 if field in ('yield_15', 'target_15') else (
            30 if field in ('yield_30', 'target_30') else None)
        self.history_requested.emit(dict(id=row['id'], row_version=row['row_version'],
                                         preferred_rule=preferred))

    def set_busy(self, busy):
        """Không cho mở editor hoặc popup trong lúc đọc/lưu/import."""
        self._busy = bool(busy)
        self.model.editing_enabled = not busy
        self.table.setEnabled(not busy)
        self.export_button.setEnabled(
            not busy and not self._exporting and self.proxy.rowCount() > 0
        )

    @staticmethod
    def _date_text(value):
        """Đưa QDate hoặc chuỗi ngày về đúng định dạng YYYYMMDD."""
        if value is None:
            return ''
        if hasattr(value, 'toString'):
            return value.toString('yyyyMMdd')
        digits = ''.join(character for character in str(value) if character.isdigit())
        if len(digits) == 8:
            return digits
        return str(value).strip()

    def set_date_range(self, date_from=None, date_to=None):
        """Chưa Search chỉ hiện tên; sau Search mới hiện khoảng From - To."""
        from_text = self._date_text(date_from)
        to_text = self._date_text(date_to)
        if from_text and to_text:
            self.title_label.setText(
                f'List of Alarm | {from_text} - {to_text}'
            )
        else:
            self.title_label.setText('List of Alarm')

    def load_rows(self, rows):
        """Hiển thị snapshot mới; giữ bộ lọc cột đang chọn."""
        self.model.set_rows(rows)
        self.proxy.invalidateFilter()
        self.update_count()

    def update_count(self, *_args):
        """Hiển thị số dòng sau lọc và tổng số alarm trong khoảng ngày."""
        self.status_label.setText(
            f'Alarm: {self.proxy.rowCount()} / {len(self.model.rows)}'
        )
        self.export_button.setEnabled(
            not self._busy and not self._exporting and self.proxy.rowCount() > 0
        )
        self.dashboard_rows_changed.emit(self.visible_rows())

    def visible_rows(self):
        """Trả snapshot đúng các dòng đang hiển thị sau bộ lọc cột Alarm."""
        rows = []
        for proxy_row in range(self.proxy.rowCount()):
            source_index = self.proxy.mapToSource(self.proxy.index(proxy_row, 0))
            if source_index.isValid():
                rows.append(self.model.rows[source_index.row()])
        return rows

    def clear_filters(self):
        """Trả lại toàn bộ alarm đã tải trong khoảng ngày."""
        self.proxy.clear_filters()
        self.update_count()

    def open_filter(self, column):
        """Popup dưới tiêu đề cột, click áp dụng ngay và chỉ có Close."""
        if not self.model.editing_enabled:
            return
        values = sorted({str(self.model.data(self.model.index(i, column)))
                         for i in range(self.model.rowCount())})
        selected = self.proxy.filters.get(column, set(values))
        dialog = FilterValueDialog(COLUMNS[column][0], values, selected,
                                   apply_immediately=True, compact=True, parent=self)
        dialog.selection_changed.connect(lambda selected: self.apply_filter(column, selected, values))
        header = self.table.horizontalHeader()
        point = header.mapToGlobal(QPoint(header.sectionViewportPosition(column), header.height()))
        screen = QApplication.screenAt(point) or QApplication.primaryScreen()
        rect = screen.availableGeometry()
        point.setX(max(rect.left(), min(point.x(), rect.right() - dialog.width())))
        point.setY(max(rect.top(), min(point.y(), rect.bottom() - dialog.height())))
        dialog.move(point)
        dialog.exec_()

    def apply_filter(self, column, selected, values):
        """Áp dụng lựa chọn và cập nhật bộ đếm."""
        self.proxy.set_filter(column, selected, values)
        self.update_count()

    def export_alarm_excel(self):
        """Chụp và xuất đúng các dòng hiện thấy sau tất cả bộ lọc cột."""
        if self._busy or self._exporting:
            return
        if self.proxy.rowCount() == 0:
            QMessageBox.information(self, 'Export Excel', 'Không có alarm để xuất.')
            return

        visible_columns = [column for column in range(self.proxy.columnCount())
                           if not self.table.isColumnHidden(column)]
        rows = [[self.proxy.data(self.proxy.index(row, column), Qt.DisplayRole)
                 for column in visible_columns]
                for row in range(self.proxy.rowCount())]
        headers = [COLUMNS[column][0] for column in visible_columns]

        alarm_dates = sorted({self.proxy.data(self.proxy.index(row, 1), Qt.EditRole)
                              for row in range(self.proxy.rowCount())})
        date_from = str(alarm_dates[0]).replace('/', '-') if alarm_dates else ''
        date_to = str(alarm_dates[-1]).replace('/', '-') if alarm_dates else ''
        default_name = f'List of Alarm from {date_from} - {date_to}.xlsx'

        output_path, _ = QFileDialog.getSaveFileName(
            self, 'Export danh sách Alarm', default_name, 'Excel Workbook (*.xlsx)')
        if not output_path:
            return
        if not output_path.lower().endswith('.xlsx'):
            output_path += '.xlsx'
            if Path(output_path).exists():
                answer = QMessageBox.question(
                    self, 'Export Excel',
                    'File đã tồn tại. Bạn có muốn ghi đè không?',
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
                if answer != QMessageBox.Yes:
                    return

        self._exporting = True
        self.export_button.setEnabled(False)
        self.export_button.setText('Exporting...')
        self._export_worker = AlarmExportWorker(output_path, headers, rows)
        self._export_worker.signals.finished.connect(self._on_alarm_export_finished)
        QThreadPool.globalInstance().start(self._export_worker)

    def _on_alarm_export_finished(self, output_path, error_message):
        """Khôi phục nút và thông báo kết quả sau khi worker hoàn thành."""
        self._exporting = False
        self._export_worker = None
        self.export_button.setText('Export Excel')
        self.export_button.setEnabled(not self._busy and self.proxy.rowCount() > 0)
        if error_message:
            QMessageBox.warning(
                self, 'Export Excel', f'Xuất Excel thất bại:\n{error_message}')
            return
        QMessageBox.information(
            self, 'Export Excel', f'Đã xuất danh sách alarm:\n{output_path}')
