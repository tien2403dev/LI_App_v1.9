from __future__ import annotations

import os
import sqlite3
import tempfile

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from domain.prime import check_cancel


REQUIRED_COLUMNS = {
    "LOTID",
    "PRODUCT",
    "EQPID",
    "TKOUTTIME",
    "INQTY",
    "OUTQTY",
    "FAILQTY",
    "YIELD",
    "TIER",
}

STAGING_INSERT_BATCH_SIZE = 5_000


class CumExcelValidationError(Exception):
    def __init__(
        self,
        source_row: int,
        message: str,
    ):
        """Khởi tạo trạng thái và các thành phần cần cho đối tượng."""
        self.source_row = source_row

        super().__init__(
            f"Dòng Excel {source_row}: {message}"
        )


@dataclass
class CumStageResult:
    staging_database_path: Path
    business_dates: set[str]
    row_count: int
    scrap_detail_count: int


class CumExcelReader:
    """
    Đọc và validate file CUM.

    Chỉ lấy 9 cột chuẩn:
    LOTID, PRODUCT, EQPID, TKOUTTIME,
    INQTY, OUTQTY, FAILQTY, YIELD, TIER.

    Các cột khác bị bỏ qua, trừ header gồm đúng 4 chữ số:
    đó là cột scrap code.
    """

    def __init__(self, progress=None, cancel=None):
        """Lưu callback tiến độ và cờ hủy; chưa nạp openpyxl."""
        self.progress = progress or (lambda message: None)
        self.cancel = cancel

    def stage_file(
        self,
        excel_path: Path,
    ) -> CumStageResult:
        """Kiểm tra file Excel rồi đọc vào staging local; dọn staging nếu lỗi."""
        check_cancel(self.cancel)
        excel_path = Path(excel_path)

        if not excel_path.exists():
            raise FileNotFoundError(
                f"Không tìm thấy file: {excel_path}"
            )

        if excel_path.suffix.lower() != ".xlsx":
            raise ValueError(
                "CUM chỉ hỗ trợ file Excel .xlsx"
            )

        staging_path = self._create_staging_database()

        try:
            (
                business_dates,
                row_count,
                scrap_detail_count,
            ) = self._read_and_stage(
                excel_path=excel_path,
                staging_path=staging_path,
            )

            return CumStageResult(
                staging_database_path=staging_path,
                business_dates=business_dates,
                row_count=row_count,
                scrap_detail_count=scrap_detail_count,
            )

        except Exception:
            self.delete_staging_database(staging_path)
            raise

    def delete_staging_database(
        self,
        staging_path: Path,
    ) -> None:
        """Xóa database staging local sau khi dùng xong hoặc khi gặp lỗi."""
        try:
            if staging_path.exists():
                staging_path.unlink()

        except OSError:
            pass

    def _read_and_stage(
        self,
        excel_path: Path,
        staging_path: Path,
    ) -> tuple[set[str], int, int]:
        # Chỉ import openpyxl khi user bấm Import CUM.
        # Không làm app khởi động chậm.
        """Đọc tuần tự sheet active, validate và ghi batch 5000 dòng vào SQLite tạm."""
        from openpyxl import load_workbook

        connection = sqlite3.connect(staging_path)

        workbook = None

        try:
            self._create_staging_tables(connection)

            workbook = load_workbook(
                filename=excel_path,
                read_only=True,
                data_only=True,
            )

            worksheet = workbook.active

            headers = self._get_headers(worksheet)

            scrap_columns = self._get_scrap_columns(
                headers=headers
            )

            business_dates: set[str] = set()

            cum_rows: list[tuple] = []
            scrap_rows: list[tuple] = []

            row_count = 0
            scrap_detail_count = 0

            lot_id_column_index = headers["LOTID"]

            for source_row, values in enumerate(
                    worksheet.iter_rows(
                        min_row=2,
                        values_only=True,
                    ),
                    start=2,
            ):
                check_cancel(self.cancel)
                lot_id_value = (
                    values[lot_id_column_index]
                    if lot_id_column_index < len(values)
                    else None
                )

                # LOTID luôn liên tục và không thiếu giữa các dòng dữ liệu.
                # Khi gặp LOTID trống, xem đây là điểm kết thúc dữ liệu
                # và bỏ qua toàn bộ nội dung thừa phía dưới sheet.
                if (
                        lot_id_value is None
                        or str(lot_id_value).strip() == ""
                ):
                    break

                cum_record, scrap_records = (
                    self._normalize_row(
                        source_row=source_row,
                        values=values,
                        headers=headers,
                        scrap_columns=scrap_columns,
                    )
                )

                business_dates.add(cum_record[1])

                cum_rows.append(cum_record)
                scrap_rows.extend(scrap_records)

                row_count += 1
                scrap_detail_count += len(scrap_records)

                if len(cum_rows) >= STAGING_INSERT_BATCH_SIZE:
                    self._insert_staging_rows(
                        connection=connection,
                        cum_rows=cum_rows,
                        scrap_rows=scrap_rows,
                    )

                    self.progress(f"Đã kiểm tra {row_count:,} dòng CUM")
                    cum_rows.clear()
                    scrap_rows.clear()

            if cum_rows:
                self._insert_staging_rows(
                    connection=connection,
                    cum_rows=cum_rows,
                    scrap_rows=scrap_rows,
                )

            if row_count == 0:
                raise ValueError(
                    "File Excel không có dòng dữ liệu nào"
                )

            check_cancel(self.cancel)
            connection.execute("CREATE INDEX idx_staging_cum_date ON staging_cum_data(DATE)")
            connection.commit()

            return (
                business_dates,
                row_count,
                scrap_detail_count,
            )

        except Exception:
            connection.rollback()
            raise

        finally:
            if workbook is not None:
                workbook.close()

            connection.close()

    def _get_headers(
        self,
        worksheet,
    ) -> dict[str, int]:
        """Đọc hàng tiêu đề và kiểm tra đủ các cột CUM bắt buộc."""
        header_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            None,
        )

        if header_row is None:
            raise ValueError(
                "File Excel không có header"
            )

        headers: dict[str, int] = {}

        for index, value in enumerate(header_row):
            if value is None:
                continue

            header_name = str(value).strip()

            if header_name:
                headers[header_name] = index

        missing_columns = REQUIRED_COLUMNS - set(headers)

        if missing_columns:
            missing_text = ", ".join(
                sorted(missing_columns)
            )

            raise ValueError(
                "Thiếu cột bắt buộc: "
                f"{missing_text}"
            )

        return headers

    @staticmethod
    def _get_scrap_columns(
        headers: dict[str, int],
    ) -> list[tuple[str, int]]:
        """
        Chỉ header đúng 4 chữ số mới được xem là scrap code.

        Các cột còn lại trong Excel CUM bị bỏ qua.
        """

        scrap_columns = []

        for header_name, column_index in headers.items():
            if (
                len(header_name) == 4
                and header_name.isdigit()
            ):
                scrap_columns.append(
                    (
                        header_name,
                        column_index,
                    )
                )

        # Giữ nguyên thứ tự cột trong Excel.
        scrap_columns.sort(
            key=lambda item: item[1]
        )

        return scrap_columns

    def _normalize_row(
        self,
        source_row: int,
        values: tuple,
        headers: dict[str, int],
        scrap_columns: list[tuple[str, int]],
    ) -> tuple[tuple, list[tuple]]:
        """Chuẩn hóa một dòng Excel thành bản ghi CUM và các dòng scrap chi tiết."""
        lot_id = self._required_text(
            values,
            headers,
            "LOTID",
            source_row,
        )

        product = self._required_text(
            values,
            headers,
            "PRODUCT",
            source_row,
        )

        eqp_id = self._required_text(
            values,
            headers,
            "EQPID",
            source_row,
        )

        raw_time = self._required_value(
            values,
            headers,
            "TKOUTTIME",
            source_row,
        )

        in_qty = self._required_integer(
            values,
            headers,
            "INQTY",
            source_row,
        )

        out_qty = self._required_integer(
            values,
            headers,
            "OUTQTY",
            source_row,
        )

        fail_qty = self._required_integer(
            values,
            headers,
            "FAILQTY",
            source_row,
        )

        yield_value = self._required_number(
            values,
            headers,
            "YIELD",
            source_row,
        )

        tier = self._required_text(
            values,
            headers,
            "TIER",
            source_row,
        )

        if len(product) < 5:
            raise CumExcelValidationError(
                source_row,
                "PRODUCT phải có ít nhất 5 ký tự",
            )

        date_text, time_text = self._normalize_datetime(
            raw_time=raw_time,
            source_row=source_row,
        )

        scrap_records = self._read_scrap_records(
            source_row=source_row,
            values=values,
            scrap_columns=scrap_columns,
        )

        scrap_text = ",".join(
            f"{scrap_code}:{qty}"
            for scrap_code, qty in scrap_records
        )

        # stage_id dùng chính số dòng Excel.
        # Giá trị này chỉ tồn tại trong staging database.
        stage_id = source_row

        cum_record = (
            stage_id,
            date_text,
            time_text,
            lot_id,
            product,
            eqp_id,
            in_qty,
            out_qty,
            fail_qty,
            yield_value,
            scrap_text,
            product[:5],
            tier,
        )

        detail_records = [
            (
                stage_id,
                scrap_code,
                qty,
            )
            for scrap_code, qty in scrap_records
        ]

        return cum_record, detail_records

    def _normalize_datetime(
        self,
        raw_time,
        source_row: int,
    ) -> tuple[str, str]:
        """Tách TKOUTTIME hợp lệ thành DATE và TIME theo định dạng database."""
        if isinstance(raw_time, datetime):
            return (
                raw_time.strftime("%Y%m%d"),
                raw_time.strftime("%H:%M:%S"),
            )

        time_text = str(raw_time).strip()

        supported_formats = (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
        )

        for time_format in supported_formats:
            try:
                parsed_time = datetime.strptime(
                    time_text,
                    time_format,
                )

                return (
                    parsed_time.strftime("%Y%m%d"),
                    parsed_time.strftime("%H:%M:%S"),
                )

            except ValueError:
                continue

        raise CumExcelValidationError(
            source_row,
            "TKOUTTIME phải có dạng "
            "YYYY-MM-DD HH:MM:SS",
        )

    def _read_scrap_records(
        self,
        source_row: int,
        values: tuple,
        scrap_columns: list[tuple[str, int]],
    ) -> list[tuple[str, int]]:
        """Lấy số lượng scrap từ các cột mã lỗi, bỏ qua ô trống và số 0."""
        scrap_records = []

        for scrap_code, column_index in scrap_columns:
            raw_qty = self._get_value_by_index(
                values,
                column_index,
            )

            # Blank hoặc 0 thì không đưa vào SCRAP.
            if raw_qty is None:
                continue

            if isinstance(raw_qty, str):
                raw_qty = raw_qty.strip()

                if raw_qty == "":
                    continue

            qty = self._normalize_scrap_qty(
                raw_qty=raw_qty,
                source_row=source_row,
                scrap_code=scrap_code,
            )

            if qty == 0:
                continue

            scrap_records.append(
                (
                    scrap_code,
                    qty,
                )
            )

        return scrap_records

    def _normalize_scrap_qty(
        self,
        raw_qty,
        source_row: int,
        scrap_code: str,
    ) -> int:
        """Kiểm tra số lượng scrap là số nguyên từ 0 đến 99."""
        if isinstance(raw_qty, int):
            qty = raw_qty

        elif isinstance(raw_qty, float):
            if not raw_qty.is_integer():
                raise CumExcelValidationError(
                    source_row,
                    f"Scrap {scrap_code} phải là số nguyên",
                )

            qty = int(raw_qty)

        else:
            qty_text = str(raw_qty).strip()

            if not qty_text.isdigit():
                raise CumExcelValidationError(
                    source_row,
                    f"Scrap {scrap_code} phải là số nguyên",
                )

            qty = int(qty_text)

        if qty < 0 or qty > 99:
            raise CumExcelValidationError(
                source_row,
                f"Scrap {scrap_code} phải nằm trong khoảng 0 đến 99",
            )

        return qty

    @staticmethod
    def _required_value(
        values: tuple,
        headers: dict[str, int],
        column_name: str,
        source_row: int,
    ):
        """Lấy ô bắt buộc và báo dòng lỗi nếu ô bị trống."""
        value = CumExcelReader._get_value(
            values,
            headers,
            column_name,
        )

        if value is None or str(value).strip() == "":
            raise CumExcelValidationError(
                source_row,
                f"Cột {column_name} không được để trống",
            )

        return value

    @staticmethod
    def _required_text(
        values: tuple,
        headers: dict[str, int],
        column_name: str,
        source_row: int,
    ) -> str:
        """Lấy giá trị bắt buộc dưới dạng chuỗi đã bỏ khoảng trắng đầu cuối."""
        value = CumExcelReader._required_value(
            values,
            headers,
            column_name,
            source_row,
        )

        return str(value).strip()

    @staticmethod
    def _required_integer(
        values: tuple,
        headers: dict[str, int],
        column_name: str,
        source_row: int,
    ) -> int:
        """Kiểm tra và chuyển giá trị bắt buộc thành số nguyên không âm."""
        value = CumExcelReader._required_value(
            values,
            headers,
            column_name,
            source_row,
        )

        if isinstance(value, int):
            number = value

        elif isinstance(value, float):
            if not value.is_integer():
                raise CumExcelValidationError(
                    source_row,
                    f"Cột {column_name} phải là số nguyên",
                )

            number = int(value)

        else:
            value_text = str(value).strip()

            if not value_text.isdigit():
                raise CumExcelValidationError(
                    source_row,
                    f"Cột {column_name} phải là số nguyên",
                )

            number = int(value_text)

        if number < 0:
            raise CumExcelValidationError(
                source_row,
                f"Cột {column_name} không được nhỏ hơn 0",
            )

        return number

    @staticmethod
    def _required_number(
        values: tuple,
        headers: dict[str, int],
        column_name: str,
        source_row: int,
    ) -> float:
        """Kiểm tra và chuyển giá trị bắt buộc thành số thực."""
        value = CumExcelReader._required_value(
            values,
            headers,
            column_name,
            source_row,
        )

        try:
            return float(value)

        except (TypeError, ValueError):
            raise CumExcelValidationError(
                source_row,
                f"Cột {column_name} phải là số",
            )

    @staticmethod
    def _get_value(
        values: tuple,
        headers: dict[str, int],
        column_name: str,
    ):
        """Lấy giá trị ô theo tên cột trong ánh xạ header."""
        return CumExcelReader._get_value_by_index(
            values,
            headers[column_name],
        )

    @staticmethod
    def _get_value_by_index(
        values: tuple,
        column_index: int,
    ):
        """Lấy giá trị ô theo vị trí, trả None nếu vượt số cột."""
        if column_index >= len(values):
            return None

        return values[column_index]

    @staticmethod
    def _is_empty_row(
        values: tuple,
    ) -> bool:
        """Kiểm tra tất cả ô trong dòng có trống hay không."""
        return all(
            value is None
            or str(value).strip() == ""
            for value in values
        )

    @staticmethod
    def _create_staging_tables(
        connection: sqlite3.Connection,
    ) -> None:
        """Tạo bảng CUM cha và bảng scrap chi tiết trong staging."""
        connection.execute(
            """
            CREATE TABLE staging_cum_data (
                stage_id INTEGER PRIMARY KEY,
                DATE TEXT NOT NULL,
                TIME TEXT NOT NULL,
                LOTID TEXT NOT NULL,
                PRODUCT TEXT NOT NULL,
                EQPID TEXT NOT NULL,
                INQTY INTEGER NOT NULL,
                OUTQTY INTEGER NOT NULL,
                FAILQTY INTEGER NOT NULL,
                YIELD REAL NOT NULL,
                SCRAP TEXT NOT NULL,
                MODEL TEXT NOT NULL,
                TIER TEXT NOT NULL
            )
            """
        )

        connection.execute(
            """
            CREATE TABLE staging_cum_scrap_detail (
                stage_id INTEGER NOT NULL,
                scrap_code TEXT NOT NULL,
                qty INTEGER NOT NULL,
                PRIMARY KEY (stage_id, scrap_code)
            )
            """
        )

    @staticmethod
    def _insert_staging_rows(
        connection: sqlite3.Connection,
        cum_rows: list[tuple],
        scrap_rows: list[tuple],
    ) -> None:
        """Ghi một batch dữ liệu CUM và scrap chi tiết vào staging."""
        connection.executemany(
            """
            INSERT INTO staging_cum_data (
                stage_id,
                DATE,
                TIME,
                LOTID,
                PRODUCT,
                EQPID,
                INQTY,
                OUTQTY,
                FAILQTY,
                YIELD,
                SCRAP,
                MODEL,
                TIER
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            cum_rows,
        )

        if scrap_rows:
            connection.executemany(
                """
                INSERT INTO staging_cum_scrap_detail (
                    stage_id,
                    scrap_code,
                    qty
                )
                VALUES (?, ?, ?)
                """,
                scrap_rows,
            )

    @staticmethod
    def _create_staging_database() -> Path:
        """Tạo file SQLite tạm trên máy cục bộ và trả đường dẫn."""
        file_descriptor, file_path = tempfile.mkstemp(
            prefix="li_cum_",
            suffix=".db",
        )

        os.close(file_descriptor)

        return Path(file_path)


