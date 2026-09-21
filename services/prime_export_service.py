from __future__ import annotations

import os
import uuid

from dataclasses import dataclass
from pathlib import Path
from typing import Union

from database.connection import create_connection


FETCH_BATCH_SIZE = 10_000
MAX_EXCEL_ROWS = 1_048_576
MAX_DATA_ROWS_PER_SHEET = MAX_EXCEL_ROWS - 1


@dataclass(frozen=True)
class PrimeExportResult:
    output_path: Path
    exported_row_count: int
    sheet_count: int


class PrimeExportService:
    """Export PRIME theo batch bằng workbook write-only."""

    def __init__(
        self,
        database_path: Union[str, Path],
    ):
        self.database_path = Path(database_path)

    def export_dates(
        self,
        business_dates: list[str],
        output_path: Union[str, Path],
    ) -> PrimeExportResult:
        """
        Export toàn bộ cột prime_data của các ngày được chọn.

        Dòng chưa có TIER vẫn được export với ô TIER trống;
        dòng đã có TIER giữ nguyên giá trị hiện tại.

        Cursor và workbook đều chạy streaming nên RAM không tăng
        theo tổng số dòng. Khi vượt giới hạn Excel, dữ liệu tự
        tách sang sheet mới.
        """

        selected_dates = sorted(
            {
                str(value).strip()
                for value in business_dates
                if str(value).strip()
            }
        )

        if not selected_dates:
            raise ValueError(
                "Vui lòng chọn ít nhất một ngày để export."
            )

        destination = Path(output_path)

        if destination.suffix.lower() != ".xlsx":
            destination = destination.with_suffix(
                ".xlsx"
            )

        destination.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        temporary_path = destination.with_name(
            f".{destination.stem}_"
            f"{uuid.uuid4().hex}.tmp.xlsx"
        )

        connection = None

        try:
            # Chỉ import openpyxl khi thực sự export.
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import (
                Alignment,
                Font,
            )
            from openpyxl.utils import (
                get_column_letter,
            )

            connection = create_connection(
                self.database_path
            )

            connection.execute(
                "PRAGMA query_only = ON"
            )

            columns = [
                row["name"]
                for row in connection.execute(
                    "PRAGMA table_info(prime_data)"
                ).fetchall()
            ]

            if not columns:
                raise RuntimeError(
                    "Không tìm thấy cấu trúc "
                    "bảng prime_data."
                )

            if "TIER" not in columns:
                raise RuntimeError(
                    "Bảng prime_data chưa có cột TIER."
                )

            quoted_columns = ", ".join(
                f'"{column}"'
                for column in columns
            )

            workbook = Workbook(
                write_only=True
            )

            sheet_index = 0
            current_sheet = None
            current_sheet_row_count = 0
            exported_row_count = 0

            def create_sheet():
                nonlocal sheet_index
                nonlocal current_sheet
                nonlocal current_sheet_row_count

                sheet_index += 1

                title = (
                    "Prime Data"
                    if sheet_index == 1
                    else f"Prime Data {sheet_index}"
                )

                current_sheet = (
                    workbook.create_sheet(
                        title=title
                    )
                )

                header_cells = []

                for column_name in columns:
                    cell = WriteOnlyCell(
                        current_sheet,
                        value=column_name,
                    )

                    cell.font = Font(
                        bold=True
                    )

                    cell.alignment = Alignment(
                        horizontal="center"
                    )

                    header_cells.append(cell)

                current_sheet.append(
                    header_cells
                )

                for (
                    column_index,
                    column_name,
                ) in enumerate(
                    columns,
                    start=1,
                ):
                    width = max(
                        12,
                        min(
                            28,
                            len(column_name) + 4,
                        ),
                    )

                    current_sheet.column_dimensions[
                        get_column_letter(
                            column_index
                        )
                    ].width = width

                current_sheet.freeze_panes = "A2"

                last_column = get_column_letter(
                    len(columns)
                )

                current_sheet.auto_filter.ref = (
                    f"A1:{last_column}1"
                )

                current_sheet_row_count = 0

            create_sheet()

            # Query từng DATE để SQLite dùng index theo DATE,
            # không tạo câu IN rất lớn.
            for business_date in selected_dates:
                cursor = connection.execute(
                    f"""
                    SELECT {quoted_columns}
                    FROM prime_data
                    WHERE DATE = ?
                    """,
                    (business_date,),
                )

                while True:
                    rows = cursor.fetchmany(
                        FETCH_BATCH_SIZE
                    )

                    if not rows:
                        break

                    for row in rows:
                        if (
                            current_sheet_row_count
                            >= MAX_DATA_ROWS_PER_SHEET
                        ):
                            create_sheet()

                        current_sheet.append(
                            tuple(row)
                        )

                        current_sheet_row_count += 1
                        exported_row_count += 1

            workbook.save(
                temporary_path
            )

            # Chỉ thay file đích sau khi workbook
            # đã được ghi hoàn chỉnh.
            os.replace(
                temporary_path,
                destination,
            )

            return PrimeExportResult(
                output_path=destination,
                exported_row_count=(
                    exported_row_count
                ),
                sheet_count=sheet_index,
            )

        finally:
            if connection is not None:
                connection.close()

            if temporary_path.exists():
                try:
                    temporary_path.unlink()
                except OSError:
                    pass